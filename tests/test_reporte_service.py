from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlmodel import Session, select

from sistema_financiero.core.caja_service import CajaService
from sistema_financiero.core.reporte_service import ReporteService
from sistema_financiero.core.tasa_cambio_service import TasaCambioService
from sistema_financiero.core.venta_controller import VentaController
from sistema_financiero.models import Producto, ReporteVentaDetalle, Usuario
from sistema_financiero.utils import TIPO_VENTA_PESO, TIPO_VENTA_UNIDAD, hoy

pytestmark = pytest.mark.integracion


# Abre una caja en la sesion (requisito para registrar ventas).
def _abrir_caja(session: Session) -> None:
    """Abre una caja en la sesion (requisito para registrar ventas)."""
    caja_sk = CajaService(db_session=session)
    if caja_sk.obtener_caja_abierta() is None:
        usuario = Usuario(
            usuario="admin_reporte",
            contrasena="hash_falso",
            nombre_completo="Admin Reporte",
        )
        session.add(usuario)
        session.commit()
        session.refresh(usuario)
        assert usuario.id is not None
        caja_sk.abrir_caja(monto_apertura_bs=Decimal("100.00"), usuario_id=usuario.id)


# Crea un producto de prueba con tipo de venta.
def _crear_producto(
    session: Session,
    nombre: str = "PROD",
    stock: Decimal = Decimal("20"),
    stock_minimo: Decimal = Decimal("5"),
    precio_bs: Decimal = Decimal("10.00"),
    tipo_venta: str = TIPO_VENTA_UNIDAD,
) -> Producto:
    producto = Producto(
        nombre_producto=nombre,
        precio_venta_bs=precio_bs,
        precio_venta_usd=Decimal("2.00"),
        stock_actual=stock,
        stock_minimo=stock_minimo,
        tipo_venta=tipo_venta,
    )
    session.add(producto)
    session.commit()
    session.refresh(producto)
    assert producto.idproducto is not None
    return producto


# Crea tasa, caja y una venta de prueba; devuelve el controlador.
def _crear_tasa_y_venta(
    session: Session,
    producto_id: int,
    cantidad: Decimal = Decimal("2"),
    precio_bs: Decimal = Decimal("10.00"),
) -> VentaController:
    ts = TasaCambioService()
    ts.registrar(
        fecha=hoy(),
        tasa_venta=Decimal("50.00"),
        tasa_compra=Decimal("50.00"),
        activa=True,
        db_session=session,
    )
    _abrir_caja(session)
    vc = VentaController()
    vc.crear(
        productos=[{"idproducto": producto_id, "cantidad": cantidad}],
        metodo_pago={"efectivo_bs": Decimal(str(precio_bs * cantidad))},
        db_session=session,
    )
    return vc


# Registra la tasa del dia y abre una caja (UNA vez por test).
def _setup_tasa_y_caja(session: Session) -> None:
    """Registra la tasa del dia y abre una caja (UNA vez por test)."""
    ts = TasaCambioService()
    ts.registrar(
        fecha=hoy(),
        tasa_venta=Decimal("50.00"),
        tasa_compra=Decimal("50.00"),
        activa=True,
        db_session=session,
    )
    _abrir_caja(session)


# Registra una venta del producto (precio unitario 10.
def _crear_venta(session: Session, producto_id: int, cantidad: Decimal) -> VentaController:
    """Registra una venta del producto (precio unitario 10.00 Bs)."""
    vc = VentaController()
    vc.crear(
        productos=[{"idproducto": producto_id, "cantidad": cantidad}],
        metodo_pago={"efectivo_bs": Decimal(str(Decimal("10.00") * cantidad))},
        db_session=session,
    )
    return vc


# Prueba generar un reporte sin ventas.
def test_generar_reporte_sin_ventas(session: Session) -> None:
    rs = ReporteService()
    reporte = rs.generar_reporte(fecha_param=hoy(), db_session=session)

    assert reporte is not None
    assert reporte.total_ventas_bs == Decimal("0.00")
    assert reporte.total_ventas_usd == Decimal("0.00")
    assert reporte.cantidad_ventas == 0
    assert reporte.unidades_vendidas == 0
    assert reporte.peso_vendido_kg == Decimal("0.000")


# Prueba el reporte con ventas del dia.
def test_generar_reporte_con_ventas(session: Session) -> None:
    producto = _crear_producto(session, precio_bs=Decimal("10.00"))
    assert producto.idproducto is not None
    _crear_tasa_y_venta(session, producto.idproducto, cantidad=Decimal("3"))

    rs = ReporteService()
    reporte = rs.generar_reporte(fecha_param=hoy(), db_session=session)

    assert reporte.cantidad_ventas == 1
    assert reporte.total_ventas_bs == Decimal("30.00")
    assert reporte.total_ventas_usd == Decimal("0.60")
    assert reporte.unidades_vendidas == 3
    assert reporte.peso_vendido_kg == Decimal("0.000")
    assert reporte.efectivo_bs == Decimal("30.00")
    assert reporte.efectivo_usd == Decimal("0.00")
    assert reporte.tarjeta == Decimal("0.00")
    assert reporte.transferencia == Decimal("0.00")


# Prueba el reporte agrupando varias ventas.
def test_generar_reporte_con_varias_ventas(session: Session) -> None:
    producto = _crear_producto(session, stock=Decimal("50"), precio_bs=Decimal("10.00"))
    ts = TasaCambioService()
    ts.registrar(
        fecha=hoy(),
        tasa_venta=Decimal("50.00"),
        tasa_compra=Decimal("50.00"),
        activa=True,
        db_session=session,
    )
    _abrir_caja(session)
    vc = VentaController()
    vc.crear(
        productos=[{"idproducto": producto.idproducto, "cantidad": 2}],
        metodo_pago={"efectivo_bs": Decimal("20.00")},
        db_session=session,
    )
    vc.crear(
        productos=[{"idproducto": producto.idproducto, "cantidad": 3}],
        metodo_pago={"efectivo_bs": Decimal("30.00")},
        db_session=session,
    )

    rs = ReporteService()
    reporte = rs.generar_reporte(fecha_param=hoy(), db_session=session)

    assert reporte.cantidad_ventas == 2
    assert reporte.total_ventas_bs == Decimal("50.00")
    assert reporte.unidades_vendidas == 5
    assert reporte.peso_vendido_kg == Decimal("0.000")


# Una venta pagada por transferencia se consolida en el reporte.
def test_generar_reporte_con_transferencia(session: Session) -> None:
    """Una venta pagada por transferencia se consolida en el reporte."""
    producto = _crear_producto(session, precio_bs=Decimal("10.00"))
    assert producto.idproducto is not None
    _abrir_caja(session)
    ts = TasaCambioService()
    ts.registrar(
        fecha=hoy(),
        tasa_venta=Decimal("50.00"),
        tasa_compra=Decimal("50.00"),
        activa=True,
        db_session=session,
    )
    vc = VentaController()
    vc.crear(
        productos=[{"idproducto": producto.idproducto, "cantidad": 2}],
        metodo_pago={
            "efectivo_bs": Decimal("0.00"),
            "transferencia": Decimal("20.00"),
        },
        db_session=session,
    )

    rs = ReporteService()
    reporte = rs.generar_reporte(fecha_param=hoy(), db_session=session)

    assert reporte.transferencia == Decimal("20.00")
    assert reporte.efectivo_bs == Decimal("0.00")


# Prueba las alertas de stock bajo y sin stock.
def test_generar_reporte_stock_bajo_y_sin_stock(session: Session) -> None:
    _crear_producto(session, nombre="BAJO", stock=Decimal("3"), stock_minimo=Decimal("5"))
    _crear_producto(session, nombre="SIN", stock=Decimal("0"), stock_minimo=Decimal("5"))
    _crear_producto(session, nombre="NORMAL", stock=Decimal("20"), stock_minimo=Decimal("5"))

    rs = ReporteService()
    reporte = rs.generar_reporte(fecha_param=hoy(), db_session=session)

    assert reporte.productos_stock_bajo >= 1
    assert reporte.productos_sin_stock >= 1


# Prueba regenerar el reporte del dia.
def test_generar_reporte_regenerar(session: Session) -> None:
    rs = ReporteService()
    rs.generar_reporte(fecha_param=hoy(), db_session=session)
    rs.generar_reporte(fecha_param=hoy(), db_session=session)

    historial = rs.listar_por_rango(desde=hoy(), hasta=hoy(), db_session=session)
    assert len(historial) == 1


# Prueba obtener un reporte por fecha existente.
def test_obtener_por_fecha_exitoso(session: Session) -> None:
    rs = ReporteService()
    rs.generar_reporte(fecha_param=hoy(), db_session=session)

    reporte = rs.obtener_por_fecha(fecha_param=hoy(), db_session=session)
    assert reporte is not None
    assert reporte.fecha == hoy()


# Prueba que devuelve None si no hay reporte.
def test_obtener_por_fecha_inexistente(session: Session) -> None:
    rs = ReporteService()
    reporte = rs.obtener_por_fecha(fecha_param=date(2020, 1, 1), db_session=session)
    assert reporte is None


# Prueba listar reportes dentro de un rango.
def test_listar_por_rango(session: Session) -> None:
    rs = ReporteService()
    rs.generar_reporte(fecha_param=hoy(), db_session=session)

    historial = rs.listar_por_rango(
        desde=date(2020, 1, 1),
        hasta=date(2030, 12, 31),
        db_session=session,
    )
    assert len(historial) >= 1


# Prueba la exportacion a Excel con sus hojas.
def test_exportar_excel(tmp_path: Path, session: Session) -> None:
    producto = _crear_producto(session, precio_bs=Decimal("10.00"))
    assert producto.idproducto is not None
    _crear_tasa_y_venta(session, producto.idproducto, cantidad=Decimal("2"))

    rs = ReporteService()
    reporte = rs.generar_reporte(fecha_param=hoy(), db_session=session)
    assert reporte.id is not None

    ruta = tmp_path / "test_reporte.xlsx"
    archivo = rs.exportar_excel(
        reporte_id=reporte.id,
        ruta_archivo=str(ruta),
        db_session=session,
    )

    assert Path(archivo).exists()
    assert Path(archivo).stat().st_size > 0

    libro = load_workbook(archivo)
    assert libro.sheetnames == [
        "Reporte Diario",
        "Ventas del Dia",
        "Alertas de Stock",
        "Productos Vendidos (Detalle)",
    ]

    hoja_resumen = libro["Reporte Diario"]
    conceptos = [
        hoja_resumen.cell(row=fila, column=1).value for fila in range(4, hoja_resumen.max_row + 1)
    ]
    assert "Unidades Vendidas" in conceptos
    assert "Peso Vendido" in conceptos

    hoja_detalle = libro["Productos Vendidos (Detalle)"]
    assert hoja_detalle["A2"].value == "PROD"
    assert hoja_detalle["B2"].value == TIPO_VENTA_UNIDAD
    assert hoja_detalle["C2"].value == 2


# Prueba que exportar un reporte inexistente falla.
def test_exportar_excel_reporte_inexistente(session: Session) -> None:
    rs = ReporteService()
    with pytest.raises(ValueError, match="No existe el reporte"):
        rs.exportar_excel(reporte_id=9999, ruta_archivo="test.xlsx", db_session=session)


# 0.
def test_reporte_peso_no_se_trunca(session: Session) -> None:
    """0.500 kg de un producto PESO NUNCA se reporta como 0 (defecto legado)."""
    producto = _crear_producto(
        session,
        nombre="HARINA",
        precio_bs=Decimal("10.00"),
        tipo_venta=TIPO_VENTA_PESO,
    )
    assert producto.idproducto is not None
    _setup_tasa_y_caja(session)
    _crear_venta(session, producto.idproducto, Decimal("0.500"))

    rs = ReporteService()
    reporte = rs.generar_reporte(fecha_param=hoy(), db_session=session)

    assert reporte.unidades_vendidas == 0
    assert reporte.peso_vendido_kg == Decimal("0.500")


# Las unidades y el peso se acumulan en campos SEPARADOS del reporte.
def test_reporte_separa_unidades_de_peso(session: Session) -> None:
    """Las unidades y el peso se acumulan en campos SEPARADOS del reporte."""
    unidad = _crear_producto(session, nombre="JUGO", precio_bs=Decimal("10.00"))
    assert unidad.idproducto is not None
    peso = _crear_producto(
        session,
        nombre="HARINA",
        precio_bs=Decimal("10.00"),
        tipo_venta=TIPO_VENTA_PESO,
    )
    assert peso.idproducto is not None
    _setup_tasa_y_caja(session)
    _crear_venta(session, unidad.idproducto, Decimal("3"))
    _crear_venta(session, peso.idproducto, Decimal("1.500"))

    rs = ReporteService()
    reporte = rs.generar_reporte(fecha_param=hoy(), db_session=session)

    assert reporte.cantidad_ventas == 2
    assert reporte.unidades_vendidas == 3
    assert reporte.peso_vendido_kg == Decimal("1.500")


# El detalle del reporte tiene UNA fila por producto con la cantidad sumada.
def test_reporte_detalle_agrupa_por_producto(session: Session) -> None:
    """El detalle del reporte tiene UNA fila por producto con la cantidad sumada."""
    producto = _crear_producto(session, nombre="AREPA", precio_bs=Decimal("10.00"))
    assert producto.idproducto is not None
    _setup_tasa_y_caja(session)
    _crear_venta(session, producto.idproducto, Decimal("2"))
    _crear_venta(session, producto.idproducto, Decimal("3"))

    rs = ReporteService()
    reporte = rs.generar_reporte(fecha_param=hoy(), db_session=session)
    assert reporte.id is not None

    detalles = session.exec(
        select(ReporteVentaDetalle).where(ReporteVentaDetalle.reporte_id == reporte.id),
    ).all()
    assert len(detalles) == 1
    assert detalles[0].nombre_producto == "AREPA"
    assert detalles[0].tipo_venta == TIPO_VENTA_UNIDAD
    assert detalles[0].producto_id == producto.idproducto
    assert detalles[0].cantidad == Decimal("5.000")

