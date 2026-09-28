from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import selectinload
from sqlmodel import Session, select

from sistema_financiero.utils import ahora
from sistema_financiero.utils import hoy as fecha_hoy
from sistema_financiero.utils.fecha import a_local, rango_dia_utc

from ..models import (
    Producto,
    ReporteDiario,
    ReporteVentaDetalle,
    Venta,
    VentaDetalle,
    obtener_sesion,
)
from ..utils import TIPO_VENTA_UNIDAD, formatear_bs, formatear_peso_kg, formatear_usd_texto

_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

_HEADER_FILL = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")

_HEADER_FONT = Font(bold=True, color="ffffff", size=11)

_COL_ESTADO: int = 5


# ReporteService: Reporte diario, detalle de productos y exportacion.
class ReporteService:
    def generar_reporte(
        self,
        fecha_param: date | None = None,
        db_session: Session | None = None,
    ) -> ReporteDiario:
        """Genera el reporte diario consolidando ventas de la fecha indicada."""
        hoy = fecha_param or fecha_hoy()

        with obtener_sesion(db_session) as session:
            desde, hasta = rango_dia_utc(hoy)

            ventas = list(
                session.exec(
                    select(Venta)
                    .options(
                        selectinload(Venta.detalles).selectinload(VentaDetalle.producto),  # type: ignore[arg-type]
                    )
                    .where(
                        Venta.fecha_venta >= desde,
                        Venta.fecha_venta <= hasta,
                        Venta.estado == "COMPLETADA",
                    ),
                ).all(),
            )

            total_bs = Decimal("0.00")
            total_usd = Decimal("0.00")
            cantidad_ventas = len(ventas)
            unidades_vendidas = 0
            peso_vendido_kg = Decimal("0.000")
            efectivo_bs = Decimal("0.00")
            efectivo_usd = Decimal("0.00")
            tarjeta = Decimal("0.00")
            pago_movil = Decimal("0.00")
            bio_pago = Decimal("0.00")
            transferencia = Decimal("0.00")

            for v in ventas:
                total_bs += v.total_bs
                total_usd += v.total_usd
                efectivo_bs += v.efectivo_bs
                efectivo_usd += v.efectivo_usd
                tarjeta += v.tarjeta
                pago_movil += v.pago_movil
                bio_pago += v.bio_pago
                transferencia += v.transferencia

            detalle_agrupado = self._agrupar_detalles(ventas)
            for (_producto_id, _nombre, tipo), cantidad in detalle_agrupado.items():
                if tipo == TIPO_VENTA_UNIDAD:
                    unidades_vendidas += int(cantidad)
                else:
                    peso_vendido_kg += cantidad
            peso_vendido_kg = peso_vendido_kg.quantize(Decimal("0.001"))

            productos = session.exec(select(Producto)).all()
            stock_bajo = sum(1 for p in productos if 0 < p.stock_actual <= p.stock_minimo)
            sin_stock = sum(1 for p in productos if p.stock_actual == 0)

            reporte_existente = session.exec(
                select(ReporteDiario).where(ReporteDiario.fecha == hoy),
            ).first()
            if reporte_existente:
                session.delete(reporte_existente)
                session.flush()

            reporte = ReporteDiario(
                fecha=hoy,
                total_ventas_bs=total_bs,
                total_ventas_usd=total_usd,
                cantidad_ventas=cantidad_ventas,
                unidades_vendidas=unidades_vendidas,
                peso_vendido_kg=peso_vendido_kg,
                productos_stock_bajo=stock_bajo,
                productos_sin_stock=sin_stock,
                efectivo_bs=efectivo_bs,
                efectivo_usd=efectivo_usd,
                tarjeta=tarjeta,
                pago_movil=pago_movil,
                bio_pago=bio_pago,
                transferencia=transferencia,
                fecha_generacion=ahora(),
            )
            session.add(reporte)
            session.flush()
            assert reporte.id is not None
            for (producto_id, nombre, tipo), cantidad in detalle_agrupado.items():
                session.add(
                    ReporteVentaDetalle(
                        reporte_id=reporte.id,
                        producto_id=producto_id,
                        nombre_producto=nombre,
                        tipo_venta=tipo,
                        cantidad=cantidad,
                    ),
                )
            session.commit()
            session.refresh(reporte)

        return reporte

    @staticmethod
    def _agrupar_detalles(
        ventas: list[Venta],
    ) -> dict[tuple[int | None, str, str], Decimal]:
        """Agrupa las cantidades vendidas por (producto, nombre, tipo de venta)."""
        agrupado: dict[tuple[int | None, str, str], Decimal] = {}
        for v in ventas:
            for d in v.detalles:
                if d.producto:
                    nombre = d.producto.nombre_producto
                    tipo = d.producto.tipo_venta
                else:
                    nombre = "Producto sin nombre"
                    tipo = TIPO_VENTA_UNIDAD
                clave = (d.producto_id, nombre, tipo)
                agrupado[clave] = agrupado.get(clave, Decimal("0.000")) + d.cantidad
        return agrupado

    def obtener_por_fecha(
        self,
        fecha_param: date | None = None,
        db_session: Session | None = None,
    ) -> ReporteDiario | None:
        """Devuelve el reporte de una fecha, o None si no existe."""
        hoy = fecha_param or fecha_hoy()
        with obtener_sesion(db_session) as session:
            return session.exec(select(ReporteDiario).where(ReporteDiario.fecha == hoy)).first()

    def listar_por_rango(
        self,
        desde: date,
        hasta: date,
        db_session: Session | None = None,
    ) -> list[ReporteDiario]:
        """Devuelve todos los reportes entre dos fechas (ordenados descendente)."""
        with obtener_sesion(db_session) as session:
            stmt = (
                select(ReporteDiario)
                .where(
                    ReporteDiario.fecha >= desde,
                    ReporteDiario.fecha <= hasta,
                )
                .order_by(ReporteDiario.fecha.desc())  # type: ignore[attr-defined]
            )
            return list(session.exec(stmt).all())

    def exportar_excel(
        self,
        reporte_id: int,
        ruta_archivo: str,
        db_session: Session | None = None,
    ) -> str:
        """Exporta un reporte a Excel. Devuelve la ruta del archivo."""
        reporte_opt, ventas, stock_bajo, sin_stock, detalles = self._cargar_datos_reporte(
            reporte_id,
            db_session,
        )
        if reporte_opt is None:
            msg = f"No existe el reporte con ID {reporte_id}"
            raise ValueError(msg)
        reporte = reporte_opt

        wb = Workbook()

        self._exportar_hoja_resumen(wb, reporte)
        self._exportar_hoja_ventas(wb, ventas)
        self._exportar_hoja_stock(wb, stock_bajo, sin_stock)
        self._exportar_hoja_detalle(wb, detalles)

        wb.save(ruta_archivo)
        return ruta_archivo

    def _cargar_datos_reporte(
        self,
        reporte_id: int,
        db_session: Session | None = None,
    ) -> tuple[
        ReporteDiario | None,
        list[Venta],
        list[Producto],
        list[Producto],
        list[ReporteVentaDetalle],
    ]:
        """Carga reporte, ventas, stock bajo, sin stock y el detalle del reporte."""
        with obtener_sesion(db_session) as session:
            reporte = session.get(ReporteDiario, reporte_id)
            if not reporte:
                return (None, [], [], [], [])

            desde, hasta = rango_dia_utc(reporte.fecha)

            ventas = list(
                session.exec(
                    select(Venta).where(
                        Venta.fecha_venta >= desde,
                        Venta.fecha_venta <= hasta,
                        Venta.estado == "COMPLETADA",
                    ),
                ).all(),
            )

            todos = session.exec(select(Producto)).all()
            stock_bajo = [p for p in todos if 0 < p.stock_actual <= p.stock_minimo]
            sin_stock = [p for p in todos if p.stock_actual == 0]

            detalles = list(
                session.exec(
                    select(ReporteVentaDetalle)
                    .where(ReporteVentaDetalle.reporte_id == reporte_id)
                    .order_by(ReporteVentaDetalle.nombre_producto),
                ).all(),
            )

        return (reporte, ventas, stock_bajo, sin_stock, detalles)

    def _exportar_hoja_resumen(self, wb: Workbook, reporte: ReporteDiario) -> None:
        """Escribe la hoja 'Reporte Diario' con el resumen del dia."""
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet("Reporte Diario", 0)
        else:
            ws.title = "Reporte Diario"
        ws.page_setup.orientation = "portrait"

        page_setup_pr = ws.sheet_properties.pageSetUpPr
        if page_setup_pr is not None:
            page_setup_pr.fitToPage = True

        ws.merge_cells("A1:B1")
        celda_titulo = ws["A1"]
        celda_titulo.value = f"Reporte Diario - {reporte.fecha}"
        celda_titulo.font = Font(bold=True, size=14)
        celda_titulo.alignment = Alignment(horizontal="center")

        for col, enc in enumerate(["Concepto", "Valor"], start=1):
            celda = ws.cell(row=3, column=col, value=enc)
            celda.font = _HEADER_FONT
            celda.fill = _HEADER_FILL
            celda.alignment = Alignment(horizontal="center")
            celda.border = _THIN_BORDER

        datos_reporte = [
            ("Total Ventas Bs.", formatear_bs(reporte.total_ventas_bs)),
            ("Total Ventas USD", formatear_usd_texto(reporte.total_ventas_usd)),
            ("Cantidad de Ventas", str(reporte.cantidad_ventas)),
            ("Unidades Vendidas", str(reporte.unidades_vendidas)),
            ("Peso Vendido", formatear_peso_kg(reporte.peso_vendido_kg)),
            ("", ""),
            ("— Metodos de Pago —", ""),
            ("Efectivo Bs.", formatear_bs(reporte.efectivo_bs)),
            ("Efectivo USD", formatear_usd_texto(reporte.efectivo_usd)),
            ("Tarjeta", formatear_bs(reporte.tarjeta)),
            ("Pago Movil", formatear_bs(reporte.pago_movil)),
            ("BioPago", formatear_bs(reporte.bio_pago)),
            ("Transferencia", formatear_bs(reporte.transferencia)),
            ("", ""),
            ("— Alertas de Stock —", ""),
            ("Productos Stock Bajo", str(reporte.productos_stock_bajo)),
            ("Productos Sin Stock", str(reporte.productos_sin_stock)),
        ]

        for i, (concepto, valor) in enumerate(datos_reporte, start=4):
            c1 = ws.cell(row=i, column=1, value=concepto)
            c2 = ws.cell(row=i, column=2, value=valor)
            c1.border = _THIN_BORDER
            c2.border = _THIN_BORDER
            if concepto.startswith("—"):
                c1.font = Font(bold=True)
            elif concepto.startswith("Total"):
                c1.font = Font(bold=True)
                c2.font = Font(bold=True)

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 25

    def _exportar_hoja_ventas(self, wb: Workbook, ventas: list[Venta]) -> None:
        """Escribe la hoja 'Ventas del Dia' con el detalle de cada venta."""
        ws = wb.create_sheet("Ventas del Dia")
        ws.page_setup.orientation = "landscape"

        enc_ventas = [
            "Factura",
            "Hora",
            "Total Bs.",
            "Total USD",
            "Efec. Bs.",
            "Efec. USD",
            "Tarjeta",
            "PagoMovil",
            "BioPago",
            "Transferencia",
        ]
        for col, enc in enumerate(enc_ventas, start=1):
            celda = ws.cell(row=1, column=col, value=enc)
            celda.font = _HEADER_FONT
            celda.fill = _HEADER_FILL
            celda.alignment = Alignment(horizontal="center")
            celda.border = _THIN_BORDER

        for i, v in enumerate(ventas, start=2):
            datos = [
                v.numero_factura or "S/N",
                a_local(v.fecha_venta).strftime("%H:%M"),
                float(v.total_bs),
                float(v.total_usd),
                float(v.efectivo_bs),
                float(v.efectivo_usd),
                float(v.tarjeta),
                float(v.pago_movil),
                float(v.bio_pago),
                float(v.transferencia),
            ]
            for col, val in enumerate(datos, start=1):
                celda = ws.cell(row=i, column=col, value=val)
                celda.border = _THIN_BORDER
                if isinstance(val, float):
                    celda.number_format = "#,##0.00"

        for col, ancho in enumerate([20, 10, 12, 12, 12, 12, 12, 12, 12, 12], start=1):
            ws.column_dimensions[chr(64 + col)].width = ancho

    def _exportar_hoja_stock(
        self,
        wb: Workbook,
        stock_bajo: list[Producto],
        sin_stock: list[Producto],
    ) -> None:
        """Escribe la hoja 'Alertas de Stock'."""
        ws = wb.create_sheet("Alertas de Stock")
        ws.page_setup.orientation = "landscape"

        enc_stock = ["Producto", "Categoria", "Stock Actual", "Stock Minimo", "Estado"]
        for col, enc in enumerate(enc_stock, start=1):
            celda = ws.cell(row=1, column=col, value=enc)
            celda.font = _HEADER_FONT
            celda.fill = _HEADER_FILL
            celda.alignment = Alignment(horizontal="center")
            celda.border = _THIN_BORDER

        fila = 2
        for p in stock_bajo:
            datos: list[str | int | Decimal] = [
                p.nombre_producto,
                p.categoria.nombre if p.categoria else "",
                p.stock_actual,
                p.stock_minimo,
                "STOCK BAJO",
            ]
            for col, val in enumerate(datos, start=1):
                celda = ws.cell(row=fila, column=col, value=val)
                celda.border = _THIN_BORDER
            fila += 1

        for p in sin_stock:
            datos = [
                p.nombre_producto,
                p.categoria.nombre if p.categoria else "",
                p.stock_actual,
                p.stock_minimo,
                "SIN STOCK",
            ]
            for col, val in enumerate(datos, start=1):
                celda = ws.cell(row=fila, column=col, value=val)
                celda.border = _THIN_BORDER
                if col == _COL_ESTADO:
                    celda.font = Font(color="ff0000", bold=True)
            fila += 1

        for letra, ancho in [("A", 30), ("B", 20), ("C", 15), ("D", 15), ("E", 15)]:
            ws.column_dimensions[letra].width = ancho

    def _exportar_hoja_detalle(
        self,
        wb: Workbook,
        detalles: list[ReporteVentaDetalle],
    ) -> None:
        """Escribe la hoja 'Productos Vendidos (Detalle)'."""
        ws = wb.create_sheet("Productos Vendidos (Detalle)")
        ws.page_setup.orientation = "landscape"

        enc_detalle = ["Producto", "Tipo Venta", "Cantidad"]
        col_cantidad = len(enc_detalle)
        for col, enc in enumerate(enc_detalle, start=1):
            celda = ws.cell(row=1, column=col, value=enc)
            celda.font = _HEADER_FONT
            celda.fill = _HEADER_FILL
            celda.alignment = Alignment(horizontal="center")
            celda.border = _THIN_BORDER

        for i, det in enumerate(detalles, start=2):
            if det.tipo_venta == TIPO_VENTA_UNIDAD:
                valor_cantidad: int | float = int(det.cantidad)
            else:
                valor_cantidad = float(det.cantidad)
            datos: list[str | int | float] = [
                det.nombre_producto,
                det.tipo_venta,
                valor_cantidad,
            ]
            for col, val in enumerate(datos, start=1):
                celda = ws.cell(row=i, column=col, value=val)
                celda.border = _THIN_BORDER
                if col == col_cantidad:
                    celda.number_format = (
                        "#,##0" if det.tipo_venta == TIPO_VENTA_UNIDAD else "#,##0.000"
                    )

        for letra, ancho in [("A", 35), ("B", 15), ("C", 15)]:
            ws.column_dimensions[letra].width = ancho

